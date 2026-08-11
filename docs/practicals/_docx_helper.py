"""Хелпер для генерации Word-документов практических работ.

Применяет стандартное оформление студенческой работы:
- Times New Roman 14 pt, межстрочный 1.5
- Поля: левое 30 мм, правое 15 мм, верх/низ 20 мм
- Абзацный отступ 1.25 см, выравнивание по ширине

Запуск:
    python3.12 _docx_helper.py  # самотест: создаёт _test.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Mm, Pt, RGBColor


FONT_NAME = "Times New Roman"
FONT_SIZE_PT = 14
FONT_SIZE_SMALL_PT = 12
LINE_SPACING = 1.5
INDENT_CM = 1.25

# Реквизиты вуза/команды
UNIVERSITY = (
    "МИНИСТЕРСТВО СЕЛЬСКОГО ХОЗЯЙСТВА РОССИЙСКОЙ ФЕДЕРАЦИИ\n"
    "Федеральное государственное бюджетное образовательное учреждение\n"
    "высшего образования\n"
    "«Российский государственный аграрный университет —\n"
    "МСХА имени К.А. Тимирязева»\n"
    "(ФГБОУ ВО РГАУ-МСХА имени К.А. Тимирязева)"
)

INSTITUTE = "Институт экономики и управления АПК"
DEPARTMENT = "Кафедра прикладной информатики"
DISCIPLINE = "Средства работы в команде"

PROJECT_NAME = (
    "Проектирование Telegram-бота и веб-приложения автоматизации доступа "
    "к расписанию занятий ФГБОУ ВО РГАУ-МСХА им. К.А. Тимирязева"
)

TEAM = [
    ("[ФИО]", "Тимлид, Backend-разработчик"),
    ("Согрин Максим", "Frontend-разработчик"),
    ("Романишин Владимир", "DevOps-инженер"),
]


def _set_run_font(run, size: int = FONT_SIZE_PT, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    # Кириллица в docx требует ручной установки шрифта
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT_NAME)


def _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, indent: bool = True,
                          space_before: float = 0, space_after: float = 0) -> None:
    pf = p.paragraph_format
    pf.alignment = alignment
    pf.line_spacing = LINE_SPACING
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if indent:
        pf.first_line_indent = Cm(INDENT_CM)
    else:
        pf.first_line_indent = Cm(0)


def new_document() -> Document:
    doc = Document()

    # Поля страницы по ГОСТ
    section = doc.sections[0]
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)

    # Стиль Normal
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(FONT_SIZE_PT)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT_NAME)

    return doc


def add_title_page(doc: Document, work_number: int, work_title: str,
                   student_role: str = "Студент") -> None:
    """Добавить титульный лист для практической работы."""

    # Шапка вуза
    for line in UNIVERSITY.split("\n"):
        p = doc.add_paragraph()
        _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
        run = p.add_run(line)
        _set_run_font(run, size=FONT_SIZE_SMALL_PT)

    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    run = p.add_run(INSTITUTE)
    _set_run_font(run, size=FONT_SIZE_SMALL_PT)

    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_before=12)
    run = p.add_run(DEPARTMENT)
    _set_run_font(run, size=FONT_SIZE_SMALL_PT, bold=True)

    # Отступ перед заголовком
    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    run = p.add_run(f"ПРАКТИЧЕСКАЯ РАБОТА № {work_number}")
    _set_run_font(run, size=16, bold=True)

    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_before=6)
    run = p.add_run(f"по дисциплине «{DISCIPLINE}»")
    _set_run_font(run, size=FONT_SIZE_PT)

    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_before=18)
    run = p.add_run(f"на тему: «{work_title}»")
    _set_run_font(run, size=FONT_SIZE_PT, bold=True)

    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_before=6)
    run = p.add_run(f"Проект: «{PROJECT_NAME}»")
    _set_run_font(run, size=FONT_SIZE_SMALL_PT)

    for _ in range(6):
        doc.add_paragraph()

    # Блок «Выполнили / Проверил»
    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
    run = p.add_run("Выполнили студенты группы [группа]:")
    _set_run_font(run)

    for member, role in TEAM:
        p = doc.add_paragraph()
        _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
        run = p.add_run(f"{member} — {role}")
        _set_run_font(run)

    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.RIGHT, indent=False, space_before=12)
    run = p.add_run("Проверил:")
    _set_run_font(run)
    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
    run = p.add_run("[научный руководитель]")
    _set_run_font(run)

    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    run = p.add_run("Москва, 2026")
    _set_run_font(run)

    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    """Заголовок уровня 1 (раздел) или 2 (подраздел)."""

    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, indent=False,
                          space_before=12, space_after=6)
    run = p.add_run(text)
    size = FONT_SIZE_PT if level == 1 else FONT_SIZE_PT
    _set_run_font(run, size=size, bold=True)


def add_centered_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                          space_before=12, space_after=6)
    run = p.add_run(text)
    _set_run_font(run, bold=True)


def add_paragraph(doc: Document, text: str, indent: bool = True,
                  alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, bold: bool = False) -> None:
    p = doc.add_paragraph()
    _set_paragraph_format(p, alignment=alignment, indent=indent)
    run = p.add_run(text)
    _set_run_font(run, bold=bold)


def add_bullet(doc: Document, text: str) -> None:
    """Маркированный пункт списка."""

    p = doc.add_paragraph(style="List Bullet")
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=False)
    run = p.runs[0] if p.runs else p.add_run("")
    run.text = text
    _set_run_font(run)


def add_numbered(doc: Document, text: str) -> None:
    """Нумерованный пункт списка."""

    p = doc.add_paragraph(style="List Number")
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=False)
    run = p.runs[0] if p.runs else p.add_run("")
    run.text = text
    _set_run_font(run)


def _set_cell_borders(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "000000")
        tcBorders.append(b)
    tcPr.append(tcBorders)


def add_table(doc: Document, headers: list[str], rows: list[list[str]],
              col_widths_cm: list[float] | None = None,
              header_bold: bool = True) -> None:
    """Таблица с границами; первая строка — заголовки."""

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Заголовки
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h)
        _set_run_font(run, size=FONT_SIZE_SMALL_PT, bold=header_bold)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_borders(cell)

    # Строки
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            # Поддержка многострочного текста
            for k, line in enumerate(str(val).split("\n")):
                if k == 0:
                    p = cell.paragraphs[0]
                else:
                    p = cell.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(line)
                _set_run_font(run, size=FONT_SIZE_SMALL_PT)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            _set_cell_borders(cell)

    # Ширина столбцов
    if col_widths_cm:
        for row in table.rows:
            for i, w in enumerate(col_widths_cm):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)

    # Промежуток после таблицы
    doc.add_paragraph()


def add_image_placeholder(doc: Document, caption: str) -> None:
    """Прямоугольник-заглушка под рисунок с подписью (когда схема рисуется в Visio/Draw.io)."""

    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = ""
    _set_cell_borders(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n\n[ Место для рисунка / схемы ]\n\n")
    _set_run_font(run, size=FONT_SIZE_SMALL_PT)

    cap = doc.add_paragraph()
    _set_paragraph_format(cap, alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                          space_before=6, space_after=12)
    run = cap.add_run(caption)
    _set_run_font(run, size=FONT_SIZE_SMALL_PT)


def save(doc: Document, path: str | Path) -> None:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


if __name__ == "__main__":
    doc = new_document()
    add_title_page(doc, 0, "Тестовый прогон")
    add_heading(doc, "1. Тестовый раздел")
    add_paragraph(doc, "Проверка оформления абзаца.")
    add_table(doc, ["Колонка А", "Колонка Б"], [["1", "2"], ["3", "4"]],
              col_widths_cm=[5, 5])
    save(doc, Path(__file__).parent / "_test.docx")
    print("OK: _test.docx")
